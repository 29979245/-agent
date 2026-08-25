"""试卷导出服务测试（6.1-6.3）：Word 双版本排版、HTML 报告与 PDF 中文。"""
import io

from docx import Document
from docx.shared import RGBColor

from app.services.question.export import (
    export_paper_docx,
    generate_paper_html,
    generate_report_html,
    register_sim_sun,
    render_paper_pdf,
    report_to_pdf,
)

QUESTIONS = [
    {
        "question_id": 1,
        "content": "写出水的化学式 H_2O，配平：Fe^{3+} + Fe → Fe^{2+}",
        "options": ["A. H_2O", "B. CO_2"],
        "answer": "A",
        "analysis": "水分子由两个氢原子和一个氧原子构成。",
        "knowledge_points": ["化学式"],
        "question_type": "选择题",
    },
    {
        "question_id": 2,
        "content": "计算 2H_2 + O_2 → 2H_2O 中生成水的物质的量。",
        "answer": "2 mol",
        "analysis": "按化学方程式系数配平计算。",
        "knowledge_points": ["物质的量", "配平"],
        "question_type": "计算题",
    },
]

STATS = {
    "total_students": 30,
    "attended": 28,
    "question_count": 2,
    "wrong_counts": {1: 12, 2: 5},
}
STUDENTS = [
    {"student_id": 1, "name": "张三", "accuracy": 0.5},
    {"student_id": 2, "name": "李四", "accuracy": 1.0},
]


def _docx_text(data):
    return "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)


# ---- 6.1 Word 排版 ----

def test_docx_sections_sealing_line(tmp_path):
    p = tmp_path / "paper.docx"
    p.write_bytes(export_paper_docx(QUESTIONS, title="期中化学"))
    text = _docx_text(p.read_bytes())
    assert "期中化学" in text
    assert "一、选择题" in text
    assert "一、计算题" in text
    assert "密 封 线" in text


def test_docx_subscript_superscript():
    data = export_paper_docx([QUESTIONS[0]])
    doc = Document(io.BytesIO(data))
    sub = [r.text for p in doc.paragraphs for r in p.runs if r.font.subscript]
    sup = [r.text for p in doc.paragraphs for r in p.runs if r.font.superscript]
    assert "2" in sub  # H_2O 的 2 为下标
    assert "3+" in sup  # Fe^{3+} 的 3+ 为上标


def test_docx_options_indented():
    data = export_paper_docx([QUESTIONS[0]])
    doc = Document(io.BytesIO(data))
    option_paras = [p for p in doc.paragraphs if p.text.startswith("A.")]
    assert option_paras and option_paras[0].paragraph_format.left_indent is not None


def test_calculation_question_leaves_answer_space():
    data = export_paper_docx([QUESTIONS[1]])
    doc = Document(io.BytesIO(data))
    assert len(doc.paragraphs) >= 4  # 标题/密封线/题干 + 答题空白


# ---- 6.2 双版本 ----

def test_student_version_hides_answers():
    text = _docx_text(export_paper_docx(QUESTIONS, with_answers=False))
    assert "（含答案版）" not in text
    assert "【答案】" not in text
    assert "【解析】" not in text


def test_teacher_version_colors_answer_and_analysis():
    data = export_paper_docx(QUESTIONS, with_answers=True)
    text = _docx_text(data)
    assert "（含答案版）" in text
    assert text.strip().endswith("（含答案版）")  # 设计 §10.1：底部标记
    assert "【答案】A" in text
    assert "【解析】" in text
    doc = Document(io.BytesIO(data))
    colors = [r.font.color.rgb for p in doc.paragraphs for r in p.runs if r.font.color and r.font.color.rgb]
    assert RGBColor(0xFF, 0x00, 0x00) in colors  # 红答案
    assert RGBColor(0x00, 0x80, 0x00) in colors  # 绿解析


# ---- 6.3 HTML 报告 + PDF ----

def test_report_html_has_cards_top5_and_kp():
    html = generate_report_html("期中化学", STATS, QUESTIONS, STUDENTS, include_analysis=True)
    assert "<meta charset=\"utf-8\">" in html
    assert "期中化学" in html
    assert "平均分" in html and "参考人数" in html and "题目数" in html
    assert "TOP5 高频错题" in html
    assert "知识点错误分布" in html
    assert "化学式" in html  # 知识点错误分布含"化学式"
    assert "75.0" in html  # 平均分 (0.5+1.0)/2*100


def test_report_html_student_version_omits_analysis():
    html = generate_report_html("期中化学", STATS, QUESTIONS, STUDENTS, include_analysis=False)
    assert "知识点错误分布" not in html
    assert "（含答案版）" not in html


def test_pdf_conversion_with_chinese(tmp_path):
    assert register_sim_sun() is True  # Windows 本机存在 SimSun
    html = generate_report_html("期中化学", STATS, QUESTIONS, STUDENTS, include_analysis=True)
    out = tmp_path / "report.pdf"
    assert report_to_pdf(html, str(out)) is True
    assert out.stat().st_size > 0
    from pypdf import PdfReader

    page = PdfReader(str(out)).pages[0]
    fonts = {str(f.get("/BaseFont")) for f in page["/Resources"].get("/Font", {}).values()}
    assert any("STSong" in name or "SimSun" in name for name in fonts)  # 中文 CID 字体已嵌入
    text = page.extract_text()
    assert "平均分" in text and "期中化学" in text  # 中文可提取、不乱码


def test_paper_html_and_pdf(tmp_path):
    """试卷 PDF（导出端点 format=pdf 依赖）：HTML 含密封线/分节/底部标记，PDF 中文可提取。"""
    html = generate_paper_html(QUESTIONS, title="期中化学", with_answers=True)
    assert "密 封 线" in html
    assert "一、选择题" in html and "一、计算题" in html
    assert "（含答案版）" in html
    assert "【答案】" in html
    pdf = render_paper_pdf(QUESTIONS, title="期中化学", with_answers=True)
    assert pdf is not None and len(pdf) > 0
    from pypdf import PdfReader

    text = PdfReader(io.BytesIO(pdf)).pages[0].extract_text()
    assert "期中化学" in text and "选择题" in text  # 中文不乱码
    assert "（含答案版）" in text

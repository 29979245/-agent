"""试卷导出服务（doc 47，设计 §6）：Word 试卷双版本 + HTML 报告转 PDF。

- export_paper_docx：python-docx 生成 A4 试卷（密封线/题型分节/化学式下标富文本）。
- with_answers 双版本：学生版无答案；教师版红答案+绿解析+「（含答案版）」。
- generate_report_html / report_to_pdf：统计报告 HTML 转 PDF，SimSun 注册防中文乱码。
"""
from __future__ import annotations

import html
import io
import logging
import os
import re
from collections import Counter

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

logger = logging.getLogger(__name__)

FONT = "SimSun"
BODY_PT = 11
TITLE_PT = 16
RED = RGBColor(0xFF, 0x00, 0x00)
GREEN = RGBColor(0x00, 0x80, 0x00)

QUESTION_TYPES = ["选择题", "填空题", "计算题", "实验题", "推断题"]

# `_2` / `_{...}` → 下标；`^2` / `^{...}` → 上标（复杂公式保留 KaTeX 源文本）
_FORMULA_TOKEN = re.compile(r"(_\{[^}]*\}|_.)|(\^\{[^}]*\}|\^.)")

_SIMSUN_CANDIDATES = [
    r"C:\Windows\Fonts\simsun.ttc",
    r"C:\Windows\Fonts\simsun.ttf",
    "/usr/share/fonts/truetype/arphic/uming.ttc",
]


# ---- Word 试卷 ----

def _set_font(run, size=BODY_PT, bold=False, color=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rpr.get_or_add_rFonts().set(qn("w:eastAsia"), FONT)


def _add_text_runs(paragraph, text, size=BODY_PT, bold=False, color=None):
    """把 `_下标`/`^上标` 拆成富文本 run；其余为普通文本 run。"""
    pos = 0
    for m in _FORMULA_TOKEN.finditer(text):
        if m.start() > pos:
            _set_font(paragraph.add_run(text[pos:m.start()]), size, bold, color)
        inner = m.group(0)[1:]
        if inner.startswith("{") and inner.endswith("}"):
            inner = inner[1:-1]
        run = paragraph.add_run(inner)
        _set_font(run, size, bold, color)
        run.font.subscript = m.group(0)[0] == "_"
        run.font.superscript = m.group(0)[0] == "^"
        pos = m.end()
    if pos < len(text):
        _set_font(paragraph.add_run(text[pos:]), size, bold, color)


def _infer_type(q: dict) -> str:
    t = q.get("question_type")
    if t in QUESTION_TYPES:
        return t
    return "选择题" if q.get("options") else "填空题"


def _add_sealing_line(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_text_runs(p, "密 封 线 （班级______ 姓名______ 学号______ 密封线内不得答题）", BODY_PT)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "dashed")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    pBdr.append(bottom)
    pPr.append(pBdr)


def _add_question(doc, num, q, qtype, with_answers):
    p = doc.add_paragraph()
    _add_text_runs(p, f"{num}. {q.get('content', '')}", BODY_PT)
    for opt in q.get("options") or []:
        op = doc.add_paragraph()
        op.paragraph_format.left_indent = Cm(1.0)
        _add_text_runs(op, opt, BODY_PT)
    if with_answers:
        if q.get("answer"):
            _add_text_runs(doc.add_paragraph(), f"【答案】{q['answer']}", BODY_PT, color=RED)
        if q.get("analysis"):
            _add_text_runs(doc.add_paragraph(), f"【解析】{q['analysis']}", BODY_PT, color=GREEN)
    # 计算题/实验题/推断题留答题空白
    if qtype in {"计算题", "实验题", "推断题"}:
        for _ in range(3):
            doc.add_paragraph()


def _add_question_sections(doc, questions, with_answers):
    by_type: dict[str, list] = {}
    for i, q in enumerate(questions, start=1):
        qtype = _infer_type(q)
        by_type.setdefault(qtype, []).append((i, q, qtype))
    for t in QUESTION_TYPES:
        items = by_type.get(t)
        if not items:
            continue
        heading = doc.add_paragraph()
        _add_text_runs(heading, f"一、{t}", BODY_PT + 2, bold=True)
        for num, q, qtype in items:
            _add_question(doc, num, q, qtype, with_answers)


def export_paper_docx(questions: list[dict], title: str = "化学试卷", with_answers: bool = False) -> bytes:
    """生成 A4 试卷 docx 字节流。questions 为 `_question_dict` 同构 dict 列表。"""
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.0)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_text_runs(p, title, TITLE_PT, bold=True)

    _add_sealing_line(doc)
    _add_question_sections(doc, questions, with_answers)

    if with_answers:
        # 设计 §10.1：教师版底部标注「（含答案版）」
        foot = doc.add_paragraph()
        foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_text_runs(foot, "（含答案版）", BODY_PT, bold=True)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---- HTML 报告 + PDF ----

def generate_report_html(
    exam_name: str,
    stats: dict,
    questions: list[dict],
    students: list[dict],
    include_analysis: bool = True,
) -> str:
    """渲染成绩报告 HTML：标题 + 统计卡 + TOP5 高频错题表；教师版含知识点错误分布。"""
    name = html.escape(exam_name)
    wrong = stats.get("wrong_counts") or {}
    qmap = {int(q.get("question_id")): q for q in questions if q.get("question_id") is not None}

    avg_score = 0.0
    if students:
        avg_score = round(sum(s.get("accuracy", 0) for s in students) / len(students) * 100, 1)

    top_rows = ""
    for qid, cnt in sorted(wrong.items(), key=lambda kv: kv[1], reverse=True)[:5]:
        q = qmap.get(int(qid)) or {}
        top_rows += (
            "<tr><td>{}</td><td>{}</td><td>{}</td><td>{}</td></tr>".format(
                qid,
                html.escape(str(q.get("content", ""))),
                html.escape("、".join(q.get("knowledge_points") or [])),
                cnt,
            )
        )

    kp_rows = ""
    if include_analysis:
        counter: Counter[str] = Counter()
        for qid, cnt in wrong.items():
            q = qmap.get(int(qid)) or {}
            for kp in q.get("knowledge_points") or []:
                counter[kp] += cnt
        kp_rows = "".join(f"<tr><td>{html.escape(kp)}</td><td>{cnt}</td></tr>" for kp, cnt in counter.most_common())

    kp_section = ""
    if include_analysis:
        kp_section = (
            "<h2>知识点错误分布</h2>"
            "<table><tr><th>知识点</th><th>错题人次</th></tr>" + kp_rows + "</table>"
        )

    return f"""<!DOCTYPE html>
<html lang="zh-CN"><head><meta charset="utf-8">
<title>{name} 成绩报告</title>
<style>
@page {{ size: A4; margin: 2cm; }}
body {{ font-family: SimSun, stsong-light, sans-serif; font-size: 11pt; color: #000; }}
h1 {{ text-align: center; font-size: 16pt; }}
h2 {{ font-size: 12pt; }}
.cards {{ display: table; width: 100%; margin: 12px 0; }}
.card {{ display: table-cell; text-align: center; border: 1px solid #000; padding: 8px; }}
.card b {{ display: block; }}
table {{ border-collapse: collapse; width: 100%; margin-bottom: 12px; }}
th, td {{ border: 1px solid #000; padding: 4px 6px; }}
th {{ background: #eee; }}
</style></head><body>
<h1>{name} 成绩报告{ '（含答案版）' if include_analysis else '' }</h1>
<div class="cards">
  <div class="card"><b>平均分</b>{avg_score}</div>
  <div class="card"><b>参考人数</b>{stats.get('attended', 0)}</div>
  <div class="card"><b>题目数</b>{stats.get('question_count', 0)}</div>
</div>
<h2>TOP5 高频错题</h2>
<table><tr><th>题号</th><th>题目内容</th><th>知识点</th><th>错题人数</th></tr>{top_rows}</table>
{kp_section}
</body></html>"""


def register_sim_sun() -> bool:
    """注册 SimSun TTF 供 reportlab/xhtml2pdf 使用，防中文乱码；返回是否注册成功。"""
    try:
        from reportlab.lib.fonts import addMapping
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
    except ImportError:
        return False
    if "SimSun" in pdfmetrics.getRegisteredFontNames():
        return True
    for path in _SIMSUN_CANDIDATES:
        if os.path.exists(path):
            pdfmetrics.registerFont(TTFont("SimSun", path))
            addMapping("SimSun", 0, 0, "SimSun")
            addMapping("SimSun", 0, 1, "SimSun")
            logger.info("[export] SimSun 字体已注册：%s", path)
            return True
    return False


def pdf_bytes(html_str: str) -> bytes | None:
    """HTML → PDF 字节流；xhtml2pdf 缺失或转换失败返回 None。"""
    try:
        from xhtml2pdf import pisa
    except ImportError:
        logger.warning("[export] xhtml2pdf 未安装，无法转 PDF")
        return None
    register_sim_sun()
    buf = io.BytesIO()
    try:
        result = pisa.CreatePDF(html_str, dest=buf, encoding="utf-8")
    except Exception as exc:  # noqa: BLE001 —— 转换失败降级，保留 HTML
        logger.warning("[export] PDF 转换失败：%s", exc)
        return None
    data = buf.getvalue()
    if result.err or not data:
        return None
    return data


def report_to_pdf(html_str: str, output_path: str) -> bool:
    """HTML 报告转 PDF（xhtml2pdf）；库缺失或转换失败返回 False。"""
    data = pdf_bytes(html_str)
    if data is None:
        return False
    try:
        with open(output_path, "wb") as f:
            f.write(data)
        return os.path.getsize(output_path) > 0
    except Exception as exc:  # noqa: BLE001 —— 写盘失败
        logger.warning("[export] PDF 写盘失败：%s", exc)
        return False


def generate_paper_html(
    questions: list[dict], title: str = "化学试卷", with_answers: bool = False
) -> str:
    """试卷 HTML（与 Word 同构：标题/密封线/题型分节/可选答案解析），供转 PDF。"""
    name = html.escape(title)
    seal = html.escape("密 封 线 （班级______ 姓名______ 学号______ 密封线内不得答题）")
    by_type: dict[str, list] = {}
    for i, q in enumerate(questions, start=1):
        by_type.setdefault(_infer_type(q), []).append((i, q))
    sections = ""
    for t in QUESTION_TYPES:
        items = by_type.get(t)
        if not items:
            continue
        rows = ""
        for num, q in items:
            row = f"<p><b>{num}.</b> {html.escape(q.get('content', ''))}</p>"
            for opt in q.get("options") or []:
                row += f"<p class=\"opt\">{html.escape(opt)}</p>"
            if with_answers:
                if q.get("answer"):
                    row += f"<p class=\"ans\">【答案】{html.escape(q['answer'])}</p>"
                if q.get("analysis"):
                    row += f"<p class=\"ana\">【解析】{html.escape(q['analysis'])}</p>"
            rows += row
        sections += f"<h2>一、{t}</h2>{rows}"
    footer = "<p class=\"foot\">（含答案版）</p>" if with_answers else ""
    return f"""<!DOCTYPE html>
<html lang="zh-CN"><head><meta charset="utf-8">
<title>{name}</title>
<style>
@page {{ size: A4; margin: 2.5cm 2cm; }}
body {{ font-family: SimSun, stsong-light, sans-serif; font-size: 11pt; color: #000; }}
h1 {{ text-align: center; font-size: 16pt; }}
h2 {{ font-size: 13pt; }}
.seal {{ text-align: center; border-bottom: 1px dashed #000; padding-bottom: 6px; }}
.opt {{ margin-left: 1cm; }}
.ans {{ color: #c00000; }}
.ana {{ color: #008000; }}
.foot {{ text-align: center; margin-top: 12px; }}
</style></head><body>
<h1>{name}</h1>
<p class="seal">{seal}</p>
{sections}
{footer}
</body></html>"""


def render_paper_pdf(
    questions: list[dict], title: str = "化学试卷", with_answers: bool = False
) -> bytes | None:
    """试卷 HTML → PDF 字节流；失败返回 None。"""
    return pdf_bytes(generate_paper_html(questions, title, with_answers))
